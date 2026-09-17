"""
Mise en forme Excel (style maison), utilisée pour écrire chaque onglet du
classeur de sortie : titre, sous-titre, en-tête coloré, bordures, formats de
nombre, largeurs de colonnes et transformation en tableau filtrable.
"""
import re
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# Couleurs de la charte
GRIS = "D9D9D9"; BLEU = "1F4E78"; VERT = "548235"; BLANC = "FFFFFF"; NOIR = "000000"

FMT_NB = "0;-0;"          # nombres entiers, les zéros restent masqués
FMT_PCT = "0.0%;-0.0%;"   # pourcentages, les zéros restent masqués

_side = Side(style="thin", color="BFBFBF")
BORD = Border(left=_side, right=_side, top=_side, bottom=_side)

TABLE_ROW = 4   # ligne de l'en-tête ; le tableau démarre donc à startrow=3


def _tname(name):
    """Transforme un nom d'onglet en identifiant de tableau Excel valide."""
    n = re.sub(r"[^A-Za-z0-9_]", "_", name)
    return n if n[:1].isalpha() else "T_" + n


def write_sheet(writer, df, sheet_name, title, subtitle="",
                header="gris", text_cols=None, pct_cols=None, min_width=0):
    """
    Écrit un DataFrame stylé dans une feuille.
    header    : "gris" (en-tête gris/bleu) ou "vert" (feuille de données brutes).
    text_cols : colonnes mises en gras noir (libellés importants).
    pct_cols  : colonnes affichées au format pourcentage.
    min_width : largeur minimale de toutes les colonnes de la feuille.
    """
    text_cols = set(text_cols or []); pct_cols = set(pct_cols or [])
    df.to_excel(writer, sheet_name=sheet_name, startrow=TABLE_ROW - 1, index=False)
    ws = writer.sheets[sheet_name]
    n_rows, n_cols = df.shape
    first, last = TABLE_ROW + 1, TABLE_ROW + n_rows

    # confort de lecture : pas de quadrillage, en-tête figé, police uniforme
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{first}"
    for row in ws.iter_rows(min_row=1, max_row=last, max_col=n_cols):
        for c in row:
            c.font = Font(name="Arial", size=10)

    # titre en ligne 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    ws.cell(1, 1, title).font = Font(name="Arial", size=16, bold=True, color=BLEU)
    ws.row_dimensions[1].height = 24

    # sous-titre en ligne 2, avec retour à la ligne automatique
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
        c = ws.cell(2, 1, subtitle)
        c.font = Font(name="Arial", size=10, italic=True, color=BLEU)
        c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        # la hauteur de la ligne s'adapte à la longueur du texte
        largeur_tot = sum(min(max(len(str(col)) + 3, min_width), 45) for col in df.columns) or 60
        n_lignes = max(1, int(len(subtitle) / max(largeur_tot, 20)) + 1)
        ws.row_dimensions[2].height = 15 * n_lignes

    # en-tête du tableau (ligne 4)
    if header == "vert":
        fill = PatternFill("solid", fgColor=VERT); font = Font(name="Arial", size=10, bold=True, color=BLANC)
    else:
        fill = PatternFill("solid", fgColor=GRIS); font = Font(name="Arial", size=10, bold=True, color=BLEU)
    for j in range(1, n_cols + 1):
        c = ws.cell(TABLE_ROW, j)
        c.fill = fill; c.font = font
        c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BORD

    # corps du tableau : bordures, formats de nombre, libellés en gras
    for j, col in enumerate(df.columns, start=1):
        for i in range(first, last + 1):
            c = ws.cell(i, j); c.border = BORD
            if col in pct_cols:
                c.number_format = FMT_PCT; c.alignment = Alignment(horizontal="center")
            elif pd.api.types.is_numeric_dtype(df[col]):
                c.number_format = FMT_NB; c.alignment = Alignment(horizontal="right")
            if col in text_cols:
                c.font = Font(name="Arial", size=10, bold=True, color=NOIR)

    # largeur des colonnes : ajustée au contenu, plafonnée à 45, plancher = min_width
    for j in range(1, n_cols + 1):
        w = max([len(str(df.columns[j - 1]))] +
                [len(str(v)) for v in df.iloc[:, j - 1].head(200)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(w + 3, min_width), 45)

    # transformation en tableau Excel : ajoute les flèches de filtre sans
    # écraser les couleurs déjà posées (style neutre)
    if n_rows:
        tbl = Table(displayName=_tname(sheet_name),
                    ref=f"A{TABLE_ROW}:{get_column_letter(n_cols)}{last}")
        tbl.tableStyleInfo = TableStyleInfo(name=None, showRowStripes=False,
                                            showColumnStripes=False)
        ws.add_table(tbl)
    return ws